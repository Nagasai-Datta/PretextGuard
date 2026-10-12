"""Rebuilds the Word copy of the master document: docs/PretextGuard_Master_Document.docx from docs/master_document.md.

    python -m pip install python-docx        (once; pandoc is also needed: brew install pandoc)
    python docs/tools/build_docx.py          (from the project root)

The Markdown file is the source and always wins; the Word file is for readers who prefer Word. The build does four things:
  1. prepares a copy of the Markdown: the one HTML table (Project snapshot) is converted by pandoc into a table pandoc can write to Word, the title block gets
     paragraph styles, and an italic caption line that repeats a figure's caption is dropped (pandoc writes the caption itself);
  2. makes a reference document (fonts, heading colours, A4 page, margins, a footer with the version and page numbers) so pandoc writes a styled file;
  3. runs pandoc;
  4. styles every table: borders, a shaded header row that repeats on each page, banded rows, small type and column widths that follow the length of the text.
Nothing here reads the data or the code of the project; it only turns one document into another. Pass --project PATH to use another project folder.
"""

import argparse
import math
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

try:
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Emu, Inches, Pt, RGBColor
except ImportError:
    sys.exit("python-docx is missing: run  python -m pip install python-docx")

NAVY, TEAL, GREY = RGBColor(0x1F, 0x2A, 0x5C), RGBColor(0x1C, 0x72, 0x93), RGBColor(0x55, 0x5B, 0x66)
TEXT_TWIPS = int((8.27 - 1.6) * 1440)          # A4 text width with 0.8 inch margins
MIN_TWIPS = int(0.62 * 1440)


# Word is strict about the order of child elements; the properties added by hand are put in the order of the schema.
ORDER = {
    "pPr": ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
            "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc",
            "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange"],
    "tblPr": ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar",
              "tblLook", "tblCaption", "tblDescription"],
    "tcPr": ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark"],
}


def reorder(element, kind):
    order = ORDER[kind]

    def key(child):
        name = child.tag.split("}")[-1]
        return order.index(name) if name in order else len(order)

    children = list(element)
    for child in children:
        element.remove(child)
    for child in sorted(children, key=key):          # sorted() is stable: unknown elements keep their order at the end
        element.append(child)


def reorder_all(root):
    for kind in ORDER:
        for element in root.iter(qn("w:" + kind)):
            reorder(element, kind)


# ------------------------------------------------------------------------------------------------ 1. the Markdown copy

def prepare_markdown(source):
    text = source.read_text(encoding="utf-8")
    start, end = text.find("<table>"), text.find("</table>")
    if start != -1 and end != -1:
        html = text[start:end + len("</table>")]
        grid = subprocess.run(["pandoc", "-f", "html", "-t", "markdown+grid_tables-smart", "--columns=200"], input=html, capture_output=True, text=True, check=True).stdout
        text = text[:start] + grid + text[end + len("</table>"):]
    title = "**PretextGuard**\n\nProject Master Document\n\n*Payload-free pretexting detection through claim verification*\n"
    if title in text:
        text = text.replace(title, '::: {custom-style="Title"}\nPretextGuard\n:::\n\n::: {custom-style="Subtitle"}\nProject Master Document\n:::\n\n'
                                   '::: {custom-style="Tagline"}\nPayload-free pretexting detection through claim verification\n:::\n', 1)
    lines = text.split("\n")
    out, i = [], 0
    while i < len(lines):
        match = re.match(r"^!\[(.+?)\]\(.+\)$", lines[i])
        out.append(lines[i])
        if match and i + 2 < len(lines) and lines[i + 1] == "" and lines[i + 2] == "*%s*" % match.group(1):
            i += 2                                   # drop the repeated italic caption
        i += 1
    return "\n".join(out)


# ------------------------------------------------------------------------------------------------ 2. the reference document

def set_font(style, name="Calibri", size=None, bold=None, italic=None, color=None):
    f = style.font
    f.name = name
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]
    if size:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color is not None:
        f.color.rgb = color


def shade_paragraph_style(style, fill, border_left=None):
    ppr = style.element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    ppr.append(shd)
    if border_left:
        bdr = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single"); left.set(qn("w:sz"), "24"); left.set(qn("w:space"), "8"); left.set(qn("w:color"), border_left)
        bdr.append(left)
        ppr.append(bdr)


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = "PAGE"
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, end):
        run._r.append(el)
    run.font.size = Pt(9); run.font.color.rgb = GREY


def make_reference(path, footer_text):
    default = subprocess.run(["pandoc", "--print-default-data-file", "reference.docx"], capture_output=True, check=True).stdout
    path.write_bytes(default)
    doc = Document(str(path))

    def get(name, kind=WD_STYLE_TYPE.PARAGRAPH):
        try:
            return doc.styles[name]
        except KeyError:
            style = doc.styles.add_style(name, kind)
            style.base_style = doc.styles["Normal"]
            return style

    normal = doc.styles["Normal"]
    set_font(normal, size=10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    for name in ("Body Text", "First Paragraph"):
        s = get(name); set_font(s, size=10.5)
        s.paragraph_format.space_after = Pt(6); s.paragraph_format.space_before = Pt(0); s.paragraph_format.line_spacing = 1.12
    s = get("Compact"); set_font(s, size=10.5); s.paragraph_format.space_after = Pt(2); s.paragraph_format.space_before = Pt(0)
    for name, (size, color, before, after) in {"Heading 1": (21, NAVY, 20, 10), "Heading 2": (15, TEAL, 16, 6), "Heading 3": (12.5, NAVY, 12, 4), "Heading 4": (11, TEAL, 10, 3)}.items():
        s = get(name); set_font(s, size=size, bold=True, color=color, italic=False)
        s.paragraph_format.space_before = Pt(before); s.paragraph_format.space_after = Pt(after); s.paragraph_format.keep_with_next = True
    h1 = doc.styles["Heading 1"]
    h1.paragraph_format.page_break_before = True
    bdr = OxmlElement("w:pBdr"); bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "8"); bottom.set(qn("w:space"), "4"); bottom.set(qn("w:color"), "1C7293")
    bdr.append(bottom); h1.element.get_or_add_pPr().append(bdr)
    t = get("Title"); set_font(t, size=34, bold=True, color=NAVY)
    t.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT; t.paragraph_format.space_before = Pt(36); t.paragraph_format.space_after = Pt(4)
    s = get("Subtitle"); set_font(s, size=20, color=TEAL); s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT; s.paragraph_format.space_after = Pt(4)
    s = get("Tagline"); set_font(s, size=12.5, italic=True, color=GREY); s.paragraph_format.space_after = Pt(16)
    s = get("Block Text"); set_font(s, size=10.5, color=NAVY)
    s.paragraph_format.left_indent = Inches(0.15); s.paragraph_format.space_before = Pt(6); s.paragraph_format.space_after = Pt(8)
    shade_paragraph_style(s, "EEF3F8", "1C7293")
    s = get("Source Code"); set_font(s, name="Consolas", size=8.5); s.paragraph_format.space_after = Pt(6); s.paragraph_format.line_spacing = 1.0
    shade_paragraph_style(s, "F3F4F6")
    set_font(get("Verbatim Char", WD_STYLE_TYPE.CHARACTER), name="Consolas", size=9, color=RGBColor(0x7A, 0x1F, 0x3D))
    cap = get("Image Caption"); set_font(cap, size=9.5, italic=True, color=GREY)
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER; cap.paragraph_format.space_after = Pt(10)
    for name in ("Captioned Figure", "Figure"):
        s = get(name); s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER; s.paragraph_format.keep_with_next = True
    set_font(get("Table Caption"), size=9.5, italic=True, color=GREY)

    section = doc.sections[0]
    section.page_width, section.page_height = Emu(int(8.27 * 914400)), Emu(int(11.69 * 914400))
    section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Inches(0.8)
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(footer_text + "   |   page "); r.font.size = Pt(9); r.font.color.rgb = GREY
    add_page_field(p)
    reorder_all(doc.styles.element)
    doc.save(str(path))


# ------------------------------------------------------------------------------------------------ 4. the tables

def shade_cell(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    for old in tcpr.findall(qn("w:shd")):
        tcpr.remove(old)
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def style_table(table):
    tblpr = table._tbl.tblPr
    for tag in ("w:tblBorders", "w:tblCellMar", "w:tblW", "w:tblLayout"):
        for old in tblpr.findall(qn(tag)):
            tblpr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4"); el.set(qn("w:space"), "0"); el.set(qn("w:color"), "B8C2CC")
        borders.append(el)
    tblpr.append(borders)
    margins = OxmlElement("w:tblCellMar")
    for edge, width in (("top", 40), ("left", 70), ("bottom", 40), ("right", 70)):
        el = OxmlElement("w:" + edge); el.set(qn("w:w"), str(width)); el.set(qn("w:type"), "dxa"); margins.append(el)
    tblpr.append(margins)
    widths = column_widths(table)
    total = OxmlElement("w:tblW"); total.set(qn("w:w"), str(sum(widths))); total.set(qn("w:type"), "dxa"); tblpr.append(total)
    layout = OxmlElement("w:tblLayout"); layout.set(qn("w:type"), "fixed"); tblpr.append(layout)
    grid = table._tbl.tblGrid
    for col in list(grid):
        grid.remove(col)
    for width in widths:
        gc = OxmlElement("w:gridCol"); gc.set(qn("w:w"), str(width)); grid.append(gc)
    for r_index, row in enumerate(table.rows):
        trpr = row._tr.get_or_add_trPr()
        if r_index == 0:
            header = OxmlElement("w:tblHeader"); header.set(qn("w:val"), "true"); trpr.append(header)
        if len(" ".join(c.text for c in row.cells)) < 600:       # a very long row may split across pages
            cant = OxmlElement("w:cantSplit"); cant.set(qn("w:val"), "true"); trpr.append(cant)
        for cell, width in zip(row.cells, widths):
            tcpr = cell._tc.get_or_add_tcPr()
            for old in tcpr.findall(qn("w:tcW")):
                tcpr.remove(old)
            tcw = OxmlElement("w:tcW"); tcw.set(qn("w:w"), str(width)); tcw.set(qn("w:type"), "dxa"); tcpr.insert(0, tcw)
            if r_index == 0:
                shade_cell(cell, "1F2A5C")
            elif r_index % 2 == 0:
                shade_cell(cell, "F3F6FA")
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(1); p.paragraph_format.space_before = Pt(0); p.paragraph_format.line_spacing = 1.0
                for run in p.runs:
                    run.font.size = Pt(8.5)
                    if r_index == 0:
                        run.font.bold = True; run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def column_widths(table):
    ncols = len(table.columns)
    weights = []
    for j in range(ncols):
        lengths = []
        for row in table.rows:
            try:
                text = row.cells[j].text
            except IndexError:
                continue
            lengths.append((len(text), max((len(w) for w in text.split()), default=1)))
        if not lengths:
            weights.append(1.0)
            continue
        ordered = sorted(l for l, _ in lengths)
        p80 = ordered[int(0.8 * (len(ordered) - 1))]
        weights.append(max(math.sqrt(max(p80, 6)), min(max(w for _, w in lengths), 24) * 0.55))
    total = sum(weights)
    widths = [max(MIN_TWIPS, int(TEXT_TWIPS * w / total)) for w in weights]
    excess, k = sum(widths) - TEXT_TWIPS, 0
    order = sorted(range(ncols), key=lambda i: -widths[i])
    while excess > 0 and k < 10000:
        col = order[k % ncols]
        cut = min(excess, max(0, widths[col] - MIN_TWIPS), 120)
        widths[col] -= cut; excess -= cut; k += 1
    widths[-1] += TEXT_TWIPS - sum(widths)
    return widths


# ------------------------------------------------------------------------------------------------ main

def main():
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    parser.add_argument("--project", default=str(Path(__file__).resolve().parents[2]))
    args = parser.parse_args()
    project = Path(args.project)
    source, figures = project / "docs" / "master_document.md", project / "docs"
    target = project / "docs" / "PretextGuard_Master_Document.docx"
    if shutil.which("pandoc") is None:
        sys.exit("pandoc is missing: brew install pandoc")
    if not source.is_file():
        sys.exit("%s does not exist" % source)
    text = source.read_text(encoding="utf-8")
    match = re.search(r"^\*\*Version:\*\* (\S+?),", text, re.M)
    version = match.group(1) if match else "?"
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        (tmp / "master_build.md").write_text(prepare_markdown(source), encoding="utf-8")
        make_reference(tmp / "reference.docx", "PretextGuard master document, version %s" % version)
        subprocess.run(["pandoc", str(tmp / "master_build.md"), "--from", "markdown+grid_tables-smart", "--resource-path=%s" % figures, "--reference-doc=%s" % (tmp / "reference.docx"),
                        "-o", str(tmp / "raw.docx")], check=True)
        doc = Document(str(tmp / "raw.docx"))
        for table in doc.tables:
            style_table(table)
        reorder_all(doc.element)
        doc.save(str(target))
        print("wrote %s (version %s, %d tables, %d kB)" % (target, version, len(doc.tables), target.stat().st_size // 1024))


if __name__ == "__main__":
    main()
